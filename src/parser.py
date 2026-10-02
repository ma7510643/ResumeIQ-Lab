from __future__ import annotations

import re
from pathlib import Path


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _from_pdf(path)
    if suffix in {".docx", ".doc"}:
        return _from_docx(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError("Please upload a PDF, DOCX, or TXT resume.")


def _from_pdf(path: Path) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = "\n".join(pages).strip()
    if not text:
        raise ValueError("Could not read text from this PDF. Scanned images are out of core scope.")
    return text


def _from_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("This DOCX file has no readable text.")
    return text


def looks_ats_friendly(text: str, filename: str) -> dict:
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    headings = [
        "education",
        "experience",
        "work experience",
        "skills",
        "projects",
        "certification",
        "summary",
        "objective",
        "achievements",
    ]
    heading_hits = sum(1 for h in headings if re.search(rf"\b{re.escape(h)}\b", text, re.I))
    word_count = len(re.findall(r"\b\w+\b", text))
    has_tables_hint = text.count("|") > 8
    two_column_hint = sum(1 for ln in lines if "  " in ln and len(ln) > 80) > 12
    issues = []
    score = 20.0
    if heading_hits < 3:
        issues.append("Add clear section headings (Education, Experience, Skills, Projects).")
        score -= 6
    if word_count < 180:
        issues.append("Resume is too short for most ATS parsers.")
        score -= 4
    if word_count > 1200:
        issues.append("Resume is long; keep it to 1–2 pages.")
        score -= 3
    if filename.lower().endswith(".pdf") is False and filename.lower().endswith(".docx") is False:
        issues.append("Prefer PDF or DOCX for ATS systems.")
        score -= 3
    if has_tables_hint or two_column_hint:
        issues.append("Complex layout detected. Single-column, simple formatting scores better in ATS.")
        score -= 4
    if re.search(r"[\u2022●■◆]", text) and text.count("•") + text.count("-") < 3:
        issues.append("Use simple bullets (- or •) instead of decorative symbols.")
        score -= 2
    return {
        "score": max(0.0, min(20.0, score)),
        "heading_hits": heading_hits,
        "word_count": word_count,
        "issues": issues,
    }
