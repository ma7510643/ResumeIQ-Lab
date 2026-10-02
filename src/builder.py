from __future__ import annotations

from io import BytesIO

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def build_resume_pdf(profile: dict) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle("Name", parent=styles["Heading1"], fontSize=18, spaceAfter=2, textColor="#0E7C7B")
    sub = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=10, textColor="#445", spaceAfter=8)
    h = ParagraphStyle("H", parent=styles["Heading2"], fontSize=12, textColor="#1B2430", spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("B", parent=styles["Normal"], fontSize=10, leading=14)

    story = []
    story.append(Paragraph(profile.get("full_name") or "Your Name", name_style))
    contact = " | ".join(
        x
        for x in [
            profile.get("email"),
            profile.get("phone"),
            profile.get("location"),
            profile.get("linkedin"),
        ]
        if x
    )
    story.append(Paragraph(contact, sub))
    if profile.get("summary"):
        story.append(Paragraph("PROFESSIONAL SUMMARY", h))
        story.append(Paragraph(profile["summary"], body))
    if profile.get("skills"):
        story.append(Paragraph("SKILLS", h))
        story.append(Paragraph(profile["skills"], body))
    if profile.get("education"):
        story.append(Paragraph("EDUCATION", h))
        for line in _lines(profile["education"]):
            story.append(Paragraph(f"• {line}", body))
    if profile.get("experience"):
        story.append(Paragraph("EXPERIENCE / INTERNSHIPS", h))
        for line in _lines(profile["experience"]):
            story.append(Paragraph(f"• {line}", body))
    if profile.get("projects"):
        story.append(Paragraph("PROJECTS", h))
        for line in _lines(profile["projects"]):
            story.append(Paragraph(f"• {line}", body))
    if profile.get("certifications"):
        story.append(Paragraph("CERTIFICATIONS", h))
        for line in _lines(profile["certifications"]):
            story.append(Paragraph(f"• {line}", body))
    story.append(Spacer(1, 8))
    doc.build(story)
    return buffer.getvalue()


def build_resume_docx(profile: dict) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(profile.get("full_name") or "Your Name", level=0)
    contact = " | ".join(
        x
        for x in [profile.get("email"), profile.get("phone"), profile.get("location"), profile.get("linkedin")]
        if x
    )
    if contact:
        doc.add_paragraph(contact)
    sections = [
        ("Professional Summary", profile.get("summary")),
        ("Skills", profile.get("skills")),
        ("Education", profile.get("education")),
        ("Experience / Internships", profile.get("experience")),
        ("Projects", profile.get("projects")),
        ("Certifications", profile.get("certifications")),
    ]
    for title, content in sections:
        if not content:
            continue
        doc.add_heading(title, level=1)
        if title in {"Skills", "Professional Summary"}:
            doc.add_paragraph(content)
        else:
            for line in _lines(content):
                doc.add_paragraph(line, style="List Bullet")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _lines(block: str) -> list[str]:
    return [ln.strip(" -•*") for ln in (block or "").splitlines() if ln.strip()]
